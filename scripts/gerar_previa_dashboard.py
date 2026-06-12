from __future__ import annotations

import collections
import datetime as dt
import os
import unicodedata
from pathlib import Path

import openpyxl
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
MONTHS = [
    "janeiro",
    "fevereiro",
    "marco",
    "abril",
    "maio",
    "junho",
    "julho",
    "agosto",
    "setembro",
    "outubro",
    "novembro",
    "dezembro",
]


def norm(value: object) -> str:
    text = str(value or "")
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return text.strip().lower()


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def parse_date(value: object) -> dt.date | None:
    if hasattr(value, "date"):
        return value.date()
    if isinstance(value, str) and value:
        return dt.datetime.fromisoformat(value).date()
    return None


def load_summary() -> dict:
    xlsx = next(path for path in ROOT.iterdir() if path.suffix.lower() == ".xlsx")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {norm(header): index for index, header in enumerate(headers)}

    monthly = collections.defaultdict(lambda: {"aprovada": 0, "rejeitada": 0, "horas": 0, "produtivas": 0, "paradas": 0})
    total_aprovado = 0
    total_rejeitado = 0
    horas_produtivas = 0.0
    horas_paradas = 0.0

    for row in ws.iter_rows(min_row=2, values_only=True):
        aprovada = row[idx["qtd aprovada"]] or 0
        rejeitada = row[idx["qtd rejeitada"]] or 0
        horas = float(row[idx["total horas"]] or 0)
        ocorrencia = str(row[idx["ocorrencia"]] or "").strip()
        date_value = parse_date(row[idx["data inicio"]])
        month = date_value.month if date_value else 0

        total_aprovado += aprovada
        total_rejeitado += rejeitada
        if ocorrencia:
            horas_paradas += horas
        else:
            horas_produtivas += horas
        if month:
            monthly[month]["aprovada"] += aprovada
            monthly[month]["rejeitada"] += rejeitada
            monthly[month]["horas"] += horas
            monthly[month]["paradas" if ocorrencia else "produtivas"] += horas

    horas_totais = horas_produtivas + horas_paradas
    return {
        "total_aprovado": total_aprovado,
        "total_rejeitado": total_rejeitado,
        "horas_produtivas": horas_produtivas,
        "horas_paradas": horas_paradas,
        "disponibilidade": horas_produtivas / horas_totais,
        "qualidade": total_aprovado / (total_aprovado + total_rejeitado),
        "monthly": monthly,
    }


def br_int(value: float) -> str:
    return f"{int(round(value)):,.0f}".replace(",", ".")


def br_percent(value: float) -> str:
    return f"{value * 100:,.2f}%".replace(",", "X").replace(".", ",").replace("X", ".")


def centered_text(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, fill: str, fnt: ImageFont.FreeTypeFont) -> None:
    bbox = draw.textbbox((0, 0), text, font=fnt)
    x = box[0] + (box[2] - box[0] - (bbox[2] - bbox[0])) / 2
    y = box[1] + (box[3] - box[1] - (bbox[3] - bbox[1])) / 2
    draw.text((x, y), text, fill=fill, font=fnt)


def draw_card(draw: ImageDraw.ImageDraw, icon_path: Path, label: str, value: str, box: tuple[int, int, int, int]) -> None:
    icon = Image.open(icon_path).convert("RGBA")
    icon.thumbnail((54, 54))
    preview.alpha_composite(icon, (box[0] + 16, box[1] + 22))
    draw.text((box[0] + 82, box[1] + 22), label, fill="#f4f2f2", font=font(18, True))
    draw.text((box[0] + 82, box[1] + 48), value, fill="#f4f2f2", font=font(31, True))


def draw_gauge(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], title: str, value: float) -> None:
    draw.text((box[0] + 16, box[1] + 16), title, fill="#f4f2f2", font=font(20, True))
    arc_box = (box[0] + 150, box[1] + 48, box[2] - 150, box[3] + 185)
    draw.arc(arc_box, 180, 360, fill="#f4f2f2", width=35)
    draw.arc(arc_box, 180, 180 + int(180 * value), fill="#45c39a", width=35)
    centered_text(draw, (box[0], box[1] + 95, box[2], box[3] - 30), br_percent(value), "#f4f2f2", font(44))
    draw.text((box[0] + 80, box[3] - 54), "0,00%", fill="#f4f2f2", font=font(16, True))
    draw.text((box[2] - 260, box[3] - 54), "100,00%", fill="#f4f2f2", font=font(16, True))


summary = load_summary()
preview = Image.open(ROOT / "Plano de Fundo.png").convert("RGBA")
preview = preview.resize((2048, 1152), Image.Resampling.LANCZOS)
draw = ImageDraw.Draw(preview)

cards = [
    ((120, 112, 430, 202), ROOT / "Ícones" / "Produzido.png", "Total Aprovado", br_int(summary["total_aprovado"])),
    ((450, 112, 760, 202), ROOT / "Ícones" / "Rejeitado.png", "Total Rejeitado", br_int(summary["total_rejeitado"])),
    ((780, 112, 1090, 202), ROOT / "Ícones" / "Hora Produtiva.png", "Horas Produtivas", br_int(summary["horas_produtivas"])),
    ((1110, 112, 1420, 202), ROOT / "Ícones" / "Hora Parada.png", "Horas Paradas", br_int(summary["horas_paradas"])),
]
for box, icon_path, label, value in cards:
    draw_card(draw, icon_path, label, value, box)

draw.text((1446, 121), "Operador", fill="#f4f2f2", font=font(17, True))
draw.text((1446, 154), "Todos", fill="#f4f2f2", font=font(20))
draw.text((1712, 121), "Mes", fill="#f4f2f2", font=font(17, True))
draw.text((1712, 154), "Todos", fill="#f4f2f2", font=font(20))

chart = (125, 250, 1900, 720)
draw.text((130, 244), "Total Aprovado por Mes", fill="#f4f2f2", font=font(20, True))
values = [summary["monthly"][i]["aprovada"] for i in range(1, 13)]
max_v = max(values)
min_v = min(values) * 0.85
points = []
for i, value in enumerate(values):
    x = chart[0] + i * ((chart[2] - chart[0]) / 11)
    y = chart[3] - ((value - min_v) / (max_v - min_v)) * (chart[3] - chart[1] - 35)
    points.append((x, y))
poly = [(chart[0], chart[3])] + points + [(chart[2], chart[3])]
draw.polygon(poly, fill=(19, 116, 111, 178))
draw.line(points, fill="#1d8d83", width=4)
for i, (x, y) in enumerate(points):
    draw.ellipse((x - 4, y - 4, x + 4, y + 4), fill="#2bb7a8")
    centered_text(draw, (int(x - 48), int(y - 42), int(x + 48), int(y - 10)), f"{round(values[i] / 1000)} Mil", "#f4f2f2", font(15, True))
    centered_text(draw, (int(x - 55), chart[3] + 6, int(x + 55), chart[3] + 36), MONTHS[i], "#f4f2f2", font(13))

draw_gauge(draw, (125, 745, 1015, 1040), "Disponibilidade", summary["disponibilidade"])
draw_gauge(draw, (1035, 745, 1900, 1040), "% Qualidade", summary["qualidade"])

out_dir = ROOT / "entregaveis"
out_dir.mkdir(exist_ok=True)
preview.save(out_dir / "preview-dashboard-producao.png")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Resumo Dashboard"
ws.append(["Indicador", "Valor"])
ws.append(["Total Aprovado", summary["total_aprovado"]])
ws.append(["Total Rejeitado", summary["total_rejeitado"]])
ws.append(["Horas Produtivas", round(summary["horas_produtivas"], 2)])
ws.append(["Horas Paradas", round(summary["horas_paradas"], 2)])
ws.append(["Disponibilidade", summary["disponibilidade"]])
ws.append(["Qualidade", summary["qualidade"]])
ws2 = wb.create_sheet("Aprovado por Mes")
ws2.append(["Mes Numero", "Mes", "Total Aprovado"])
for i, month in enumerate(MONTHS, start=1):
    ws2.append([i, month, summary["monthly"][i]["aprovada"]])
wb.save(out_dir / "resumo-dashboard-producao.xlsx")

print("Preview gerada:", out_dir / "preview-dashboard-producao.png")
print("Resumo gerado:", out_dir / "resumo-dashboard-producao.xlsx")
