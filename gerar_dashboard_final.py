from pathlib import Path
import collections
import datetime as dt
import os
import unicodedata

import openpyxl
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
MONTHS = [
    "janeiro",
    "fevereiro",
    "marco",
    "abril",
    "maio",
    "junho",
    "julho",
    "agosto",
    "setembro",
    "outubro",
    "novembro",
    "dezembro",
]


def ascii_key(value):
    text = str(value or "")
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii").strip().lower()


def font(size, bold=False):
    candidates = [
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def parse_date(value):
    if hasattr(value, "date"):
        return value.date()
    if isinstance(value, str) and value:
        return dt.datetime.fromisoformat(value).date()
    return None


def br_int(value):
    return f"{int(round(value)):,.0f}".replace(",", ".")


def br_pct(value):
    return f"{value * 100:,.2f}%".replace(",", "X").replace(".", ",").replace("X", ".")


def find_icons_dir():
    for name in os.listdir(ROOT):
        path = ROOT / name
        if path.is_dir() and "cone" in ascii_key(name):
            return path
    raise FileNotFoundError("Pasta de icones nao encontrada")


def find_icon(icons_dir, name_part):
    target = ascii_key(name_part)
    for path in icons_dir.iterdir():
        if target in ascii_key(path.name):
            return path
    raise FileNotFoundError(name_part)


def centered_text(draw, box, text, fill, text_font):
    bounds = draw.textbbox((0, 0), text, font=text_font)
    x = box[0] + (box[2] - box[0] - (bounds[2] - bounds[0])) / 2
    y = box[1] + (box[3] - box[1] - (bounds[3] - bounds[1])) / 2
    draw.text((x, y), text, fill=fill, font=text_font)


def load_summary():
    xlsx = next(name for name in os.listdir(ROOT) if name.lower().endswith(".xlsx"))
    workbook = openpyxl.load_workbook(ROOT / xlsx, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    indexes = {ascii_key(header): index for index, header in enumerate(headers)}

    monthly = collections.defaultdict(lambda: {"approved": 0})
    approved = 0
    rejected = 0
    productive_hours = 0.0
    stopped_hours = 0.0

    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_approved = row[indexes["qtd aprovada"]] or 0
        row_rejected = row[indexes["qtd rejeitada"]] or 0
        hours = float(row[indexes["total horas"]] or 0)
        occurrence = str(row[indexes["ocorrencia"]] or "").strip()
        date_value = parse_date(row[indexes["data inicio"]])

        approved += row_approved
        rejected += row_rejected
        if occurrence:
            stopped_hours += hours
        else:
            productive_hours += hours
        if date_value:
            monthly[date_value.month]["approved"] += row_approved

    return {
        "approved": approved,
        "rejected": rejected,
        "productive_hours": productive_hours,
        "stopped_hours": stopped_hours,
        "availability": productive_hours / (productive_hours + stopped_hours),
        "quality": approved / (approved + rejected),
        "monthly": monthly,
    }


summary = load_summary()
icons_dir = find_icons_dir()
image = Image.open(ROOT / "Plano de Fundo.png").convert("RGBA")
image = image.resize((1280, 720), Image.Resampling.LANCZOS)
draw = ImageDraw.Draw(image)


def draw_card(icon_path, label, value, box):
    icon = Image.open(icon_path).convert("RGBA")
    icon.thumbnail((34, 34))
    image.alpha_composite(icon, (box[0] + 10, box[1] + 14))
    draw.text((box[0] + 52, box[1] + 14), label, fill="#f4f2f2", font=font(11, True))
    draw.text((box[0] + 52, box[1] + 33), value, fill="#f4f2f2", font=font(20, True))


draw_card(find_icon(icons_dir, "Produzido"), "Total Aprovado", br_int(summary["approved"]), (80, 72, 286, 132))
draw_card(find_icon(icons_dir, "Rejeitado"), "Total Rejeitado", br_int(summary["rejected"]), (292, 72, 498, 132))
draw_card(find_icon(icons_dir, "Hora Produtiva"), "Horas Produtivas", br_int(summary["productive_hours"]), (506, 72, 712, 132))
draw_card(find_icon(icons_dir, "Hora Parada"), "Horas Paradas", br_int(summary["stopped_hours"]), (720, 72, 926, 132))

draw.text((904, 80), "Operador", fill="#f4f2f2", font=font(12, True))
draw.text((904, 104), "Todos", fill="#f4f2f2", font=font(13))
draw.text((1070, 80), "Mes", fill="#f4f2f2", font=font(12, True))
draw.text((1070, 104), "Todos", fill="#f4f2f2", font=font(13))

chart = (78, 160, 1190, 450)
draw.text((82, 157), "Total Aprovado por Mes", fill="#f4f2f2", font=font(13, True))
values = [summary["monthly"][month]["approved"] for month in range(1, 13)]
max_value = max(values)
min_value = min(values) * 0.85
points = []
for index, value in enumerate(values):
    x = chart[0] + index * ((chart[2] - chart[0]) / 11)
    y = chart[3] - ((value - min_value) / (max_value - min_value)) * (chart[3] - chart[1] - 24)
    points.append((x, y))

draw.polygon([(chart[0], chart[3])] + points + [(chart[2], chart[3])], fill=(19, 116, 111, 178))
draw.line(points, fill="#1d8d83", width=3)
for index, (x, y) in enumerate(points):
    draw.ellipse((x - 3, y - 3, x + 3, y + 3), fill="#2bb7a8")
    centered_text(draw, (int(x - 30), int(y - 28), int(x + 30), int(y - 8)), f"{round(values[index] / 1000)} Mil", "#f4f2f2", font(9, True))
    centered_text(draw, (int(x - 35), chart[3] + 4, int(x + 35), chart[3] + 24), MONTHS[index], "#f4f2f2", font(8))


def draw_gauge(box, title, value):
    draw.text((box[0] + 8, box[1] + 10), title, fill="#f4f2f2", font=font(13, True))
    arc = (box[0] + 90, box[1] + 30, box[2] - 90, box[3] + 120)
    draw.arc(arc, 180, 360, fill="#f4f2f2", width=22)
    draw.arc(arc, 180, 180 + int(180 * value), fill="#45c39a", width=22)
    centered_text(draw, (box[0], box[1] + 70, box[2], box[3] - 10), br_pct(value), "#f4f2f2", font(29))
    draw.text((box[0] + 52, box[3] - 35), "0,00%", fill="#f4f2f2", font=font(9, True))
    draw.text((box[2] - 160, box[3] - 35), "100,00%", fill="#f4f2f2", font=font(9, True))


draw_gauge((80, 492, 662, 692), "Disponibilidade", summary["availability"])
draw_gauge((670, 492, 1260, 692), "% Qualidade", summary["quality"])

(ROOT / "entregaveis").mkdir(exist_ok=True)
image.convert("RGB").save(ROOT / "Dashboard_Producao_Final.png", quality=95)
image.convert("RGB").save(ROOT / "entregaveis" / "Dashboard_Producao_Final.png", quality=95)

print("Dashboard_Producao_Final.png")
print(br_int(summary["approved"]), br_int(summary["rejected"]), br_int(summary["productive_hours"]), br_int(summary["stopped_hours"]), br_pct(summary["availability"]), br_pct(summary["quality"]))
